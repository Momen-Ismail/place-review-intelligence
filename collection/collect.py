from __future__ import annotations

import argparse
from datetime import datetime
import json
import re
from typing import Any

import pandas as pd

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from collection.collectors.apify_google import ApifyGoogleReviewsCollector
from collection.core.database import ReviewDatabase
from collection.core.normalizer import normalize_items
from config import EXPORT_DIR


# ============================================================
# GENERAL HELPERS
# ============================================================


def slugify(value: str) -> str:
    """
    Convert a place name into a safe filename.

    Example:
    Al Assiel Restaurant & Shisha
    ->
    Al_Assiel_Restaurant_Shisha
    """

    value = re.sub(
        r"[^a-zA-Z0-9_-]+",
        "_",
        value.strip(),
    )

    return value.strip("_")[:80] or "place"


def load_json_dict(value: Any) -> dict:
    """
    Safely convert JSON text into a Python dictionary.
    """

    if not value:
        return {}

    if isinstance(value, dict):
        return value

    try:
        result = json.loads(value)

        if isinstance(result, dict):
            return result

    except (json.JSONDecodeError, TypeError):
        pass

    return {}


def load_json_list(value: Any) -> list:
    """
    Safely convert JSON text into a Python list.
    """

    if not value:
        return []

    if isinstance(value, list):
        return value

    try:
        result = json.loads(value)

        if isinstance(result, list):
            return result

    except (json.JSONDecodeError, TypeError):
        pass

    return []


def split_timestamp(
    value: str | None,
) -> tuple[str, str]:
    """
    Convert:

    2026-09-03T18:17:14.642Z

    into:

    date = 2026-09-03
    time = 18:17:14
    """

    if not value:
        return "", ""

    try:
        normalized = value.replace(
            "Z",
            "+00:00",
        )

        dt = datetime.fromisoformat(
            normalized
        )

        return (
            dt.date().isoformat(),
            dt.strftime("%H:%M:%S"),
        )

    except ValueError:
        return "", ""


# ============================================================
# HUMAN-READABLE REVIEW EXPORT
# ============================================================


def build_export_rows(
    reviews: list[dict],
    place: dict,
    include_raw_json: bool = False,
) -> list[dict]:
    """
    Convert normalized review data into a clean,
    human-readable table for Excel / CSV.

    Important rating levels:

    1. Google overall rating
       = rating of the whole place

    2. Review rating
       = rating given by one reviewer

    3. Detailed ratings
       = Food / Service / Atmosphere / etc.
    """

    export_rows: list[dict] = []

    for review in reviews:

        context = load_json_dict(
            review.get(
                "review_context_json"
            )
        )

        detailed_ratings = load_json_dict(
            review.get(
                "detailed_ratings_json"
            )
        )

        images = load_json_list(
            review.get(
                "review_images_json"
            )
        )

        published_date, published_time = (
            split_timestamp(
                review.get(
                    "published_at"
                )
            )
        )

        row = {
            # =================================================
            # PLACE / BUSINESS INFORMATION
            # =================================================

            "Place name":
                place.get("title"),

            "Place category":
                place.get("category"),

            "Google overall rating":
                place.get("total_score"),

            "Google review count":
                place.get("reviews_count"),

            # =================================================
            # INDIVIDUAL REVIEW RATING
            # =================================================

            "Review rating":
                review.get("rating"),

            # =================================================
            # REVIEWER INFORMATION
            # =================================================

            "Reviewer name":
                review.get(
                    "reviewer_name"
                ),

            "Reviewer ID":
                review.get(
                    "reviewer_id"
                ),

            "Local Guide":
                review.get(
                    "is_local_guide"
                ),

            "Reviewer profile URL":
                review.get(
                    "reviewer_url"
                ),

            # =================================================
            # REVIEW TEXT
            # =================================================

            "Review text - original":
                review.get(
                    "text_original"
                ),

            "Review text - translated":
                review.get(
                    "text_translated"
                ),

            "Original language":
                review.get(
                    "original_language"
                ),

            "Translated language":
                review.get(
                    "translated_language"
                ),

            # =================================================
            # REVIEW DATE / TIME
            # =================================================

            "Published date":
                published_date,

            "Published time UTC":
                published_time,

            "Published timestamp":
                review.get(
                    "published_at"
                ),

            "Published relative":
                review.get(
                    "published_relative"
                ),

            "Visited in":
                review.get(
                    "visited_in"
                ),

            # =================================================
            # REVIEW ENGAGEMENT
            # =================================================

            "Likes":
                review.get(
                    "likes_count"
                ),

            # =================================================
            # OWNER RESPONSE
            # =================================================

            "Owner response":
                review.get(
                    "owner_response_text"
                ),

            "Owner response date":
                review.get(
                    "owner_response_date"
                ),

            # =================================================
            # PHOTOS ATTACHED TO THIS REVIEW
            # =================================================

            "Review photo count":
                len(images),

            "Review photo URLs":
                " | ".join(
                    str(image)
                    for image in images
                ),

            # =================================================
            # REFERENCES
            # =================================================

            "Review URL":
                review.get(
                    "review_url"
                ),

            "Review ID":
                review.get(
                    "review_id"
                ),

            "Place ID":
                review.get(
                    "place_id"
                ),

            "Source":
                review.get(
                    "source"
                ),

            "Scraped at":
                review.get(
                    "scraped_at"
                ),
        }

        # =====================================================
        # DETAILED CATEGORY RATINGS
        #
        # Restaurant examples:
        #
        # Food rating
        # Service rating
        # Atmosphere rating
        #
        # Other place types may provide different ratings.
        # =====================================================

        for key, value in detailed_ratings.items():

            row[
                f"{key} rating"
            ] = value

        # =====================================================
        # REVIEW CONTEXT
        #
        # Examples:
        #
        # Meal type
        # Price per person
        # Group size
        # Wait time
        # Parking space
        # Noise level
        # Recommendation for vegetarians
        # Vegetarian offerings
        #
        # These are created dynamically.
        # =====================================================

        for key, value in context.items():
            row[key] = value

        # =====================================================
        # OPTIONAL RAW APIFY RECORD
        # =====================================================

        if include_raw_json:

            row["Raw Apify JSON"] = (
                review.get(
                    "raw_json"
                )
            )

        export_rows.append(row)

    return export_rows


# ============================================================
# PLACE INFORMATION SHEET
# ============================================================


def build_place_info(
    place: dict,
) -> dict:
    """
    Create readable information about the place itself.
    """

    return {
        "Place name":
            place.get("title"),

        "Category":
            place.get("category"),

        "All categories":
            place.get(
                "categories_json"
            ),

        "Google overall rating":
            place.get(
                "total_score"
            ),

        "Google review count":
            place.get(
                "reviews_count"
            ),

        "Price":
            place.get("price"),

        "Address":
            place.get("address"),

        "Street":
            place.get("street"),

        "City":
            place.get("city"),

        "State":
            place.get("state"),

        "Country":
            place.get(
                "country_code"
            ),

        "Postal code":
            place.get(
                "postal_code"
            ),

        "Latitude":
            place.get("latitude"),

        "Longitude":
            place.get("longitude"),

        "Google Maps URL":
            place.get(
                "google_maps_url"
            ),

        "Place ID":
            place.get("place_id"),

        "CID":
            place.get("cid"),

        "Last collected":
            place.get(
                "updated_at"
            ),
    }


# ============================================================
# EXCEL FORMATTING
# ============================================================


def format_excel_sheet(
    worksheet,
):
    """
    Apply simple readable formatting to an Excel sheet.
    """

    # Keep header visible.
    worksheet.freeze_panes = "A2"

    # Enable filters.
    if worksheet.max_row >= 1:
        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

    # Header formatting.
    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    # Column widths.
    for column_cells in worksheet.columns:

        column_letter = get_column_letter(
            column_cells[0].column
        )

        header = str(
            column_cells[0].value
            or ""
        )

        max_length = len(header)

        for cell in column_cells[1:100]:

            value = cell.value

            if value is None:
                continue

            value_length = min(
                len(str(value)),
                60,
            )

            max_length = max(
                max_length,
                value_length,
            )

        if "Review text" in header:

            width = 60

        elif "Owner response" in header:

            width = 50

        elif "URL" in header:

            width = 45

        elif "ID" in header:

            width = 28

        elif "rating" in header.lower():

            width = 18

        else:

            width = min(
                max(
                    max_length + 2,
                    12,
                ),
                35,
            )

        worksheet.column_dimensions[
            column_letter
        ].width = width

    # Wrap text in all data rows.
    for row in worksheet.iter_rows(
        min_row=2
    ):

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )


# ============================================================
# EXCEL EXPORT
# ============================================================


def export_excel(
    output,
    place: dict,
    review_rows: list[dict],
):
    """
    Create an Excel workbook containing:

    1. Reviews
    2. Place Info
    """

    reviews_df = pd.DataFrame(
        review_rows
    )

    place_df = pd.DataFrame(
        [
            build_place_info(
                place
            )
        ]
    )

    with pd.ExcelWriter(
        output,
        engine="openpyxl",
    ) as writer:

        reviews_df.to_excel(
            writer,
            sheet_name="Reviews",
            index=False,
        )

        place_df.to_excel(
            writer,
            sheet_name="Place Info",
            index=False,
        )

        reviews_sheet = (
            writer.book["Reviews"]
        )

        place_sheet = (
            writer.book["Place Info"]
        )

        format_excel_sheet(
            reviews_sheet
        )

        format_excel_sheet(
            place_sheet
        )


# ============================================================
# COLLECTION COMMAND
# ============================================================


def collect_command(args):
    """
    Collect Google Maps reviews through Apify,
    normalize them,
    store them in SQLite,
    and export CSV / Excel when requested.
    """

    collector = (
        ApifyGoogleReviewsCollector()
    )

    if args.max_reviews is None:

        print(
            "Collecting all available reviews "
            "per place..."
        )

    else:

        print(
            f"Collecting up to "
            f"{args.max_reviews} review(s) "
            f"per place..."
        )

    # --------------------------------------------------------
    # APIFY COLLECTION
    # --------------------------------------------------------

    raw_items = collector.collect(
        urls=args.url,
        max_reviews=args.max_reviews,
        sort=args.sort,
        language=args.language,
        include_personal_data=(
            args.personal_data
        ),
    )

    print(
        f"Apify returned "
        f"{len(raw_items)} "
        f"raw review item(s)."
    )

    if not raw_items:

        print(
            "No reviews were returned by Apify."
        )

        return

    # --------------------------------------------------------
    # NORMALIZE
    # --------------------------------------------------------

    places, reviews = normalize_items(
        raw_items
    )

    # --------------------------------------------------------
    # DATABASE
    # --------------------------------------------------------

    db = ReviewDatabase()

    try:

        db.upsert_places(
            places
        )

        inserted, updated = (
            db.upsert_reviews(
                reviews
            )
        )

        print(
            f"Places processed: "
            f"{len(places)}"
        )

        print(
            f"New reviews inserted: "
            f"{inserted}"
        )

        print(
            f"Existing reviews refreshed: "
            f"{updated}"
        )

        # ----------------------------------------------------
        # EXPORT
        # ----------------------------------------------------

        if (
            args.export_csv
            or args.export_excel
        ):

            EXPORT_DIR.mkdir(
                parents=True,
                exist_ok=True,
            )

            for place in places:

                place_id = (
                    place["place_id"]
                )

                place_reviews = [
                    review
                    for review in reviews
                    if (
                        review["place_id"]
                        == place_id
                    )
                ]

                # IMPORTANT:
                # pass BOTH the reviews and the place.
                export_rows = build_export_rows(
                    place_reviews,
                    place,
                    include_raw_json=(
                        args.include_raw_json
                    ),
                )

                safe_name = slugify(
                    place.get("title")
                    or place_id
                )

                # --------------------------------------------
                # CSV
                # --------------------------------------------

                if args.export_csv:

                    csv_output = (
                        EXPORT_DIR
                        / f"{safe_name}_reviews.csv"
                    )

                    pd.DataFrame(
                        export_rows
                    ).to_csv(
                        csv_output,
                        index=False,
                        encoding="utf-8-sig",
                    )

                    print(
                        f"CSV export: "
                        f"{csv_output}"
                    )

                # --------------------------------------------
                # EXCEL
                # --------------------------------------------

                if args.export_excel:

                    excel_output = (
                        EXPORT_DIR
                        / f"{safe_name}_reviews.xlsx"
                    )

                    export_excel(
                        excel_output,
                        place,
                        export_rows,
                    )

                    print(
                        f"Excel export: "
                        f"{excel_output}"
                    )

    finally:

        db.close()


# ============================================================
# LIST PLACES COMMAND
# ============================================================


def places_command(_args):

    db = ReviewDatabase()

    try:

        rows = db.list_places()

        if not rows:

            print(
                "No places stored yet."
            )

            return

        for row in rows:

            print(
                f"{row['title']} | "
                f"{row['category']} | "
                f"{row['city'] or '-'} | "
                f"stored="
                f"{row['stored_reviews']} | "
                f"avg="
                f"{row['stored_avg_rating']} | "
                f"id="
                f"{row['place_id']}"
            )

    finally:

        db.close()


# ============================================================
# PLACE STATS COMMAND
# ============================================================


def stats_command(args):

    db = ReviewDatabase()

    try:

        summary, distribution = (
            db.place_stats(
                args.place_id
            )
        )

        if not summary:

            print(
                "Place not found."
            )

            return

        print(
            "\nPLACE SUMMARY"
        )

        print(
            "-" * 50
        )

        print(
            f"Name: "
            f"{summary['title']}"
        )

        print(
            f"Category: "
            f"{summary['category']}"
        )

        print(
            f"Location: "
            f"{summary['city'] or '-'}, "
            f"{summary['country_code'] or '-'}"
        )

        print(
            f"Stored reviews: "
            f"{summary['stored_reviews']}"
        )

        print(
            f"Stored average: "
            f"{summary['stored_avg_rating']}"
        )

        print(
            f"Google overall rating: "
            f"{summary['google_score']}"
        )

        print(
            f"Google review count: "
            f"{summary['google_review_count']}"
        )

        print(
            f"Stored date range: "
            f"{summary['oldest_review']} "
            f"-> "
            f"{summary['newest_review']}"
        )

        print(
            "\nRATING DISTRIBUTION"
        )

        print(
            "-" * 50
        )

        for row in distribution:

            print(
                f"{row['rating']} star: "
                f"{row['count']}"
            )

    finally:

        db.close()


# ============================================================
# COMMAND-LINE ARGUMENTS
# ============================================================


def build_parser():

    parser = argparse.ArgumentParser(
        description=(
            "Generic place review "
            "intelligence collector."
        )
    )

    sub = parser.add_subparsers(
        dest="command",
        required=True,
    )

    # --------------------------------------------------------
    # COLLECT
    # --------------------------------------------------------

    collect = sub.add_parser(
        "collect",
        help=(
            "Collect Google Maps "
            "reviews via Apify."
        ),
    )

    collect.add_argument(
        "--url",
        action="append",
        required=True,
        help=(
            "Google Maps place URL. "
            "Repeat --url for multiple places."
        ),
    )

    # None = collect all available reviews.
    collect.add_argument(
        "--max-reviews",
        type=int,
        default=None,
        help=(
            "Maximum number of reviews per place. "
            "Omit to collect all available reviews."
        ),
    )

    collect.add_argument(
        "--sort",
        default="newest",
        choices=[
            "newest",
            "mostRelevant",
            "highestRanking",
            "lowestRanking",
        ],
    )

    collect.add_argument(
        "--language",
        default="en",
        help=(
            "Google Maps interface language. "
            "Reviews in all available languages "
            "are still collected."
        ),
    )

    # Reviewer profile information is enabled.
    collect.add_argument(
        "--personal-data",
        action=argparse.BooleanOptionalAction,
        default=True,
        help=(
            "Collect reviewer Google profile data. "
            "Enabled by default. "
            "Use --no-personal-data to disable."
        ),
    )

    collect.add_argument(
        "--export-csv",
        action="store_true",
        help=(
            "Export a readable CSV file."
        ),
    )

    collect.add_argument(
        "--export-excel",
        action="store_true",
        help=(
            "Export a formatted Excel .xlsx file."
        ),
    )

    collect.add_argument(
        "--include-raw-json",
        action="store_true",
        help=(
            "Include complete raw Apify JSON "
            "inside exported files."
        ),
    )

    collect.set_defaults(
        func=collect_command
    )

    # --------------------------------------------------------
    # PLACES
    # --------------------------------------------------------

    places = sub.add_parser(
        "places",
        help=(
            "List places stored "
            "in the database."
        ),
    )

    places.set_defaults(
        func=places_command
    )

    # --------------------------------------------------------
    # STATS
    # --------------------------------------------------------

    stats = sub.add_parser(
        "stats",
        help=(
            "Show basic statistics "
            "for one stored place."
        ),
    )

    stats.add_argument(
        "place_id"
    )

    stats.set_defaults(
        func=stats_command
    )

    return parser


# ============================================================
# MAIN
# ============================================================


def main():

    parser = build_parser()

    args = parser.parse_args()

    args.func(args)


if __name__ == "__main__":

    main()
