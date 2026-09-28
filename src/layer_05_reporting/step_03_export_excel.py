"""Export the compact management workbook."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _value(value):
    return json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value


def _frame(rows):
    if not rows: return pd.DataFrame({"Message": ["No qualifying data"]})
    frame = pd.DataFrame(rows)
    frame.columns = [str(column).replace("_", " ").title() for column in frame.columns]
    return frame.map(_value)


def _format(sheet):
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True); cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for cells in sheet.columns:
        width = min(max(12, *(len(str(cell.value or "")) for cell in cells[:150])) + 2, 62)
        sheet.column_dimensions[get_column_letter(cells[0].column)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", horizontal="right" if isinstance(cell.value, (int, float)) else "left", wrap_text=True)
            if isinstance(cell.value, bool):
                cell.number_format = "General"
            elif isinstance(cell.value, int):
                cell.number_format = "#,##0"
            elif isinstance(cell.value, float):
                cell.number_format = "0.00"


def _flatten(prefix, value, rows):
    if isinstance(value, dict):
        for key, nested in value.items(): _flatten(f"{prefix}.{key}" if prefix else key, nested, rows)
    elif isinstance(value, list):
        rows.append({"field": prefix, "value": " | ".join(map(str, value))})
    else: rows.append({"field": prefix, "value": value})


def _summary(analysis):
    overall = analysis["overall_performance"]; trend = analysis["trends"]["summary"]
    strengths = analysis["satisfaction_vs_dissatisfaction"]["strengths"][:5]
    weaknesses = analysis["satisfaction_vs_dissatisfaction"]["weaknesses"][:5]
    priorities = analysis["business_priorities"]["priorities"][:5]
    mixed_topics = {row["topic"] for row in analysis["emerging_insights"].get("mixed_emerging_signals", [])}
    rows = [
        {"section": "Restaurant Overview", "metric": "Google Rating", "value": overall["google_rating"]},
        {"section": "Restaurant Overview", "metric": "Collected Average Rating", "value": overall["collected_average_rating"]},
        {"section": "Restaurant Overview", "metric": "Total Reviews", "value": overall["collected_reviews"]},
        {"section": "Restaurant Overview", "metric": "Written Reviews", "value": overall["written_reviews"]},
        {"section": "Restaurant Overview", "metric": "Written Review Coverage %", "value": overall["written_review_coverage_pct"]},
    ]
    for label, key in (("Recent 30-Day Average", "recent_30d_average"), ("Previous 30-Day Average", "previous_30d_average"),
                       ("30-Day Change", "rating_change_recent_30d_vs_previous_30d"), ("Recent 90-Day Average", "recent_90d_average"),
                       ("Previous 90-Day Average", "previous_90d_average"), ("90-Day Change", "rating_change_recent_90d_vs_previous_90d"),
                       ("Recent 90-Day vs Historical", "rating_change_recent_90d_vs_historical")):
        rows.append({"section": "Recent Performance", "metric": label, "value": trend.get(key)})
    for item in strengths:
        note = "; mixed recent signal—see Emerging Issues" if item["topic"] in mixed_topics else ""
        rows.append({"section": "Key Strengths", "metric": item["topic"], "value": f"{item['positive_reviews']} positive reviews ({item['evidence_level']} evidence){note}"})
    for item in weaknesses:
        note = "; mixed recent signal—see Emerging Issues" if item["topic"] in mixed_topics else ""
        rows.append({"section": "Key Problems", "metric": item["topic"], "value": f"{item['negative_reviews']} negative reviews ({item['evidence_level']} evidence){note}"})
    for item in priorities: rows.append({"section": "Top Business Priorities", "metric": f"{item['priority_rank']}. {item['topic']}", "value": item["reason"]})
    for item in analysis["structured_ratings"]: rows.append({"section": "Structured Ratings", "metric": item["category"], "value": item["average_rating"]})
    return rows


def _buffet_rows(value):
    rows = [{"section": "overview", "metric": key, "value": nested}
            for key, nested in value.items() if not isinstance(nested, (dict, list))]
    for key in ("overall_sentiment", "quality", "variety", "freshness", "replenishment", "value", "waiting_queue"):
        rows.append({"section": "subtopic", "segment": key, **value[key]})
    for period, metrics in value.get("period_comparison", {}).items():
        rows.append({"section": "period comparison", "segment": period, **metrics})
    return rows


def _context_rows(value):
    return value.get("values", [])


def _cadence_rows(value):
    rows = [{"section": "period", "segment": name, **metrics} for name, metrics in value["periods"].items()]
    rows += [{"section": "monthly volume", "segment": row["month"], "review_count": row["review_count"]}
             for row in value["monthly_review_volume"]]
    return rows


def _owner_response_rows(value):
    overall = value["overall"]
    rows = [{"section": "period", "segment": "All time", "reviews": overall["total_reviews"],
             "responses": overall["reviews_with_owner_response"], "no_response": overall["reviews_without_owner_response"],
             "response_rate_pct": overall["owner_response_rate_pct"],
             "low_score_reviews": value["low_score"]["low_score_reviews"],
             "low_score_responses": value["low_score"]["low_score_responses"],
             "low_score_response_rate_pct": value["low_score"]["low_score_response_rate_pct"]}]
    for name, metrics in value["periods"].items():
        rows.append({"section": "period", "segment": name, "reviews": metrics["review_count"],
                     "responses": metrics["responses"], "no_response": metrics["no_response"],
                     "response_rate_pct": metrics["response_rate_pct"],
                     "low_score_reviews": metrics["low_score_review_count"],
                     "low_score_responses": metrics["low_score_responses"],
                     "low_score_response_rate_pct": metrics["low_score_response_rate_pct"]})
    for metrics in value["by_star_rating"]:
        rows.append({"section": "by star rating", "segment": f"{metrics['rating']} star", "reviews": metrics["total_reviews"],
                     "responses": metrics["responses"], "no_response": metrics["no_response"],
                     "response_rate_pct": metrics["response_rate_pct"]})
    for metric, result in value["response_delay"].items():
        rows.append({"section": "response delay", "segment": metric, "value": result})
    return rows


def _topic_comparison_rows(value):
    rows = []
    for item in value:
        historical = item["periods"]["historical"]
        previous = item["periods"]["previous_90d"]
        recent = item["periods"]["recent_90d"]
        rows.append({
            "topic": item["topic"], "business_area": item["business_area"],
            "historical_mentions": historical["reviews_mentioning_topic"],
            "historical_negative_pct": historical["negative_share_of_topic_reviews_pct"],
            "previous_90d_mentions": previous["reviews_mentioning_topic"],
            "previous_90d_negative_pct": previous["negative_share_of_topic_reviews_pct"],
            "recent_90d_mentions": recent["reviews_mentioning_topic"],
            "recent_90d_negative_pct": recent["negative_share_of_topic_reviews_pct"],
            "recent_vs_historical_negative_change_pp": item["recent_vs_historical_negative_share_change_pp"],
            "recent_vs_previous90_negative_change_pp": item["recent_vs_previous90_negative_share_change_pp"],
            "historical_positive_pct": historical["positive_share_of_topic_reviews_pct"],
            "recent_positive_pct": recent["positive_share_of_topic_reviews_pct"],
            "recent_evidence": recent["evidence_level"],
        })
    return rows


def _low_score_summary_rows(value):
    rows = []
    for name, metrics in value["groups"].items():
        rows.append({"rating_group": name, "reviews": metrics["review_count"],
                     "share_of_all_reviews_pct": metrics["share_of_all_reviews_pct"],
                     "written_reviews": metrics["written_review_count"],
                     "owner_responses": metrics["owner_response_count"],
                     "response_rate_pct": metrics["owner_response_rate_pct"]})
    return rows


def _menu_analysis_rows(value, favourite_signals, management_insights):
    rows = [{"section": "summary", "metric": key, "value": result} for key, result in value["summary"].items()]
    rows += [{"section": "most discussed signal", "rank": rank, **item}
             for rank, item in enumerate(favourite_signals.get("most_discussed", []), 1)]
    rows += [{"section": "most positively discussed signal", "rank": rank, **item}
             for rank, item in enumerate(favourite_signals.get("most_positively_discussed", []), 1)]
    rows += [{"section": "negatively discussed item", **item}
             for item in management_insights["negatively_discussed_items"]]
    rows += [{"section": "strongest category review attention", **item}
             for item in management_insights["categories_with_strongest_review_attention"]]
    rows += [{"section": "weak or no category review attention", **item}
             for item in management_insights["categories_with_weak_or_no_review_attention"]]
    rows += [{"section": "item review attention", **item} for item in value["items"]]
    return rows


def _bar_drinks_rows(value):
    rows = []

    def add(metric, period, result, evidence):
        rows.append({"metric": metric, "period": period, "value": result, "evidence": evidence})

    overall = value["overall_review_evidence"]
    topic_evidence = overall["topic_evidence_level"]
    add("Drinks mentions", "All written reviews", overall["topic_mentions"], topic_evidence)
    add("Written-review coverage %", "All written reviews", overall["written_review_coverage_pct"], topic_evidence)
    add("Positive", "All written reviews", overall["positive"], overall["positive_evidence_level"])
    add("Negative", "All written reviews", overall["negative"], overall["negative_issue_evidence_level"])
    add("Neutral", "All written reviews", overall["neutral"], overall["neutral_evidence_level"])
    add("Average rating when mentioned", "All written reviews",
        overall["average_overall_review_rating_when_mentioned"], topic_evidence)
    for period, metrics in value["period_comparison"].items():
        evidence = metrics["topic_evidence_level"]
        add("Drinks mentions", period, metrics["mentions"], evidence)
        add("Mention coverage %", period, metrics["mention_coverage_pct_of_written_reviews"], evidence)
        add("Positive", period, metrics["positive"], metrics["positive_evidence_level"])
        add("Negative", period, metrics["negative"], metrics["negative_issue_evidence_level"])
        add("Neutral", period, metrics["neutral"], metrics["neutral_evidence_level"])
        add("Average rating when mentioned", period,
            metrics["average_overall_review_rating_when_mentioned"], evidence)
    totals = value["menu_attention"]["totals"]
    add("Drink menu items", "Current menu", totals["total_drink_items"], "menu catalog")
    add("Drink items mentioned", "Current menu + reviews",
        totals["drink_items_with_observable_review_mentions"], "review attention")
    add("Drink items with no observable review mentions", "Current menu + reviews",
        totals["drink_items_with_no_observable_review_mentions"], "review attention")
    add("Drink review-attention coverage %", "Current menu + reviews",
        totals["pct_drink_items_with_review_attention"], "review attention")
    for category in value["menu_attention"]["categories"]:
        add("Category items / items mentioned", category["category"],
            f"{category['total_menu_items']} / {category['items_with_review_mentions']}",
            category["review_attention_status"])
    add("Management note", "Interpretation", value["terminology_note"], "Review attention is not sales")
    return rows


def _journey_rows(value):
    return [{"stage": row["stage"], "customer_evidence": row["customer_evidence"],
             "management_observations": row["management_observations"], "potential_actions": row["potential_actions"]}
            for row in value["stages"]]


def export_analysis_excel(output_path: Path | str, analysis: dict, prepared: dict | None = None) -> None:
    prepared = prepared or {}
    sheets = {
        "Summary": _summary(analysis),
        "Review Cadence": _cadence_rows(analysis["review_cadence"]),
        "Owner Responses": _owner_response_rows(analysis["owner_response_analysis"]),
        "Period Comparison": _topic_comparison_rows(analysis["topic_period_comparison"]),
        "Low Score Analysis": _low_score_summary_rows(analysis["low_rating_analysis"]),
        "Low Score Topics": analysis["low_rating_analysis"]["period_comparison"]["topics"],
        "Rating Distribution": analysis["overall_performance"]["rating_distribution"],
        "Structured Ratings": analysis["structured_ratings"], "Monthly Trend": analysis["trends"]["monthly"],
        "Topics": analysis["topics"], "Dishes": analysis["dishes"],
        "Buffet Analysis": _buffet_rows(analysis["buffet_analysis"]),
        "Bar & Drinks": _bar_drinks_rows(analysis["bar_drinks_analysis"]),
        "Menu Analysis": _menu_analysis_rows(
            analysis["menu_analysis"], analysis["customer_favourite_signals"], analysis["menu_management_insights"]
        ),
        "Menu Categories": analysis["menu_management_insights"]["category_review_attention"],
        "Business Priorities": analysis["business_priorities"]["priorities"],
        "Customer Journey": _journey_rows(analysis["customer_journey"]),
        "Management Baseline": [],
        "Context": _context_rows(analysis["context"]),
        "Data Quality": [{"metric": key, "value": value} for key, value in analysis["data_quality"].items()],
    }
    _flatten("", analysis["management_baseline"], sheets["Management Baseline"])
    output_path = Path(output_path); output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for name, rows in sheets.items():
            _frame(rows).to_excel(writer, sheet_name=name[:31], index=False)
            _format(writer.book[name[:31]])
        writer.book["Summary"].column_dimensions["A"].width = 28
        writer.book["Summary"].column_dimensions["B"].width = 40
        writer.book["Summary"].column_dimensions["C"].width = 70
