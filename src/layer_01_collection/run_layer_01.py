from __future__ import annotations
import argparse
from datetime import datetime, timezone
import json
from pathlib import Path
import re
import sqlite3
import sys
from typing import Any
import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from src.layer_01_collection.step_01_fetch_google_reviews import ApifyGoogleReviewsCollector
from src.layer_01_collection.step_02_normalize_reviews import normalize_items
from src.layer_01_collection.step_03_store_reviews import ReviewDatabase
from src.layer_01_collection.step_04_export_reviews import export_paths, save_collection_summary, save_database_summary
from config import (
    BUSINESS_NAME,
    COLLECTION_SNAPSHOT_PATH,
    DATABASE_PATH,
    EXPORT_DIR,
)


def existing_review_count(path: Path = DATABASE_PATH) -> int:
    """Return the stored review count without creating or changing a database."""
    if not path.exists():
        return 0
    connection = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    try:
        row = connection.execute(
            "SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name='reviews'"
        ).fetchone()
        if not row or not row[0]:
            return 0
        return int(connection.execute("SELECT COUNT(*) FROM reviews").fetchone()[0])
    finally:
        connection.close()


def write_collection_snapshot(urls, started_at, reviews, collector) -> None:
    dated = sorted(
        str(review.get("published_at") or "")
        for review in reviews
        if review.get("published_at")
    )
    payload = {
        "restaurant": BUSINESS_NAME,
        "google_maps_url": urls[0] if len(urls) == 1 else urls,
        "collection_started_at": started_at,
        "collection_completed_at": datetime.now(timezone.utc).isoformat(),
        "reviews_collected": len(reviews),
        "latest_review_date": dated[-1][:10] if dated else None,
        "earliest_review_date": dated[0][:10] if dated else None,
        "actor_run": collector.last_run_info,
        "full_history_attempted": True,
    }
    COLLECTION_SNAPSHOT_PATH.parent.mkdir(parents=True, exist_ok=True)
    COLLECTION_SNAPSHOT_PATH.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

def slugify(value: str) -> str:
    """
    Convert a place name into a safe filename.

    Example:
    Al Assiel Restaurant & Shisha
    ->
    Al_Assiel_Restaurant_Shisha
    """
    value = re.sub('[^a-zA-Z0-9_-]+', '_', value.strip())
    return value.strip('_')[:80] or 'place'

def load_json_dict(value: Any) -> dict:
    """
    Safely convert JSON text into a Python dictionary.
    """
    if not value:
        return {}
    if isinstance(value, dict):
        return value
    try:
        result = json.loads(value)
        if isinstance(result, dict):
            return result
    except (json.JSONDecodeError, TypeError):
        pass
    return {}

def load_json_list(value: Any) -> list:
    """
    Safely convert JSON text into a Python list.
    """
    if not value:
        return []
    if isinstance(value, list):
        return value
    try:
        result = json.loads(value)
        if isinstance(result, list):
            return result
    except (json.JSONDecodeError, TypeError):
        pass
    return []

def split_timestamp(value: str | None) -> tuple[str, str]:
    """
    Convert:

    2026-09-03T18:17:14.642Z

    into:

    date = 2026-09-03
    time = 18:17:14
    """
    if not value:
        return ('', '')
    try:
        normalized = value.replace('Z', '+00:00')
        dt = datetime.fromisoformat(normalized)
        return (dt.date().isoformat(), dt.strftime('%H:%M:%S'))
    except ValueError:
        return ('', '')

def build_export_rows(reviews: list[dict], place: dict, include_raw_json: bool=False) -> list[dict]:
    """
    Convert normalized review data into a clean,
    human-readable table for Excel / CSV.

    Important rating levels:

    1. Google overall rating
       = rating of the whole place

    2. Review rating
       = rating given by one reviewer

    3. Detailed ratings
       = Food / Service / Atmosphere / etc.
    """
    export_rows: list[dict] = []
    for review in reviews:
        context = load_json_dict(review.get('review_context_json'))
        detailed_ratings = load_json_dict(review.get('detailed_ratings_json'))
        images = load_json_list(review.get('review_images_json'))
        published_date, published_time = split_timestamp(review.get('published_at'))
        row = {'Place name': place.get('title'), 'Place category': place.get('category'), 'Google overall rating': place.get('total_score'), 'Google review count': place.get('reviews_count'), 'Review rating': review.get('rating'), 'Reviewer name': review.get('reviewer_name'), 'Reviewer ID': review.get('reviewer_id'), 'Local Guide': review.get('is_local_guide'), 'Reviewer profile URL': review.get('reviewer_url'), 'Review text - original': review.get('text_original'), 'Review text - translated': review.get('text_translated'), 'Original language': review.get('original_language'), 'Translated language': review.get('translated_language'), 'Published date': published_date, 'Published time UTC': published_time, 'Published timestamp': review.get('published_at'), 'Published relative': review.get('published_relative'), 'Visited in': review.get('visited_in'), 'Likes': review.get('likes_count'), 'Owner response': review.get('owner_response_text'), 'Owner response date': review.get('owner_response_date'), 'Review photo count': len(images), 'Review photo URLs': ' | '.join((str(image) for image in images)), 'Review URL': review.get('review_url'), 'Review ID': review.get('review_id'), 'Place ID': review.get('place_id'), 'Source': review.get('source'), 'Scraped at': review.get('scraped_at')}
        for key, value in detailed_ratings.items():
            row[f'{key} rating'] = value
        for key, value in context.items():
            row[key] = value
        if include_raw_json:
            row['Raw Apify JSON'] = review.get('raw_json')
        export_rows.append(row)
    return export_rows

def build_place_info(place: dict) -> dict:
    """
    Create readable information about the place itself.
    """
    return {'Place name': place.get('title'), 'Category': place.get('category'), 'All categories': place.get('categories_json'), 'Google overall rating': place.get('total_score'), 'Google review count': place.get('reviews_count'), 'Price': place.get('price'), 'Address': place.get('address'), 'Street': place.get('street'), 'City': place.get('city'), 'State': place.get('state'), 'Country': place.get('country_code'), 'Postal code': place.get('postal_code'), 'Latitude': place.get('latitude'), 'Longitude': place.get('longitude'), 'Google Maps URL': place.get('google_maps_url'), 'Place ID': place.get('place_id'), 'CID': place.get('cid'), 'Last collected': place.get('updated_at')}

def format_excel_sheet(worksheet):
    """
    Apply simple readable formatting to an Excel sheet.
    """
    worksheet.freeze_panes = 'A2'
    if worksheet.max_row >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        header = str(column_cells[0].value or '')
        max_length = len(header)
        for cell in column_cells[1:100]:
            value = cell.value
            if value is None:
                continue
            value_length = min(len(str(value)), 60)
            max_length = max(max_length, value_length)
        if 'Review text' in header:
            width = 60
        elif 'Owner response' in header:
            width = 50
        elif 'URL' in header:
            width = 45
        elif 'ID' in header:
            width = 28
        elif 'rating' in header.lower():
            width = 18
        else:
            width = min(max(max_length + 2, 12), 35)
        worksheet.column_dimensions[column_letter].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)

def export_excel(output, place: dict, review_rows: list[dict]):
    """
    Create an Excel workbook containing:

    1. Reviews
    2. Place Info
    """
    reviews_df = pd.DataFrame(review_rows)
    place_df = pd.DataFrame([build_place_info(place)])
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        reviews_df.to_excel(writer, sheet_name='Reviews', index=False)
        place_df.to_excel(writer, sheet_name='Place Info', index=False)
        reviews_sheet = writer.book['Reviews']
        place_sheet = writer.book['Place Info']
        format_excel_sheet(reviews_sheet)
        format_excel_sheet(place_sheet)

def collect_command(args):
    """
    Collect Google Maps reviews through Apify,
    normalize them,
    store them in SQLite,
    and export CSV / Excel when requested.
    """
    stored = existing_review_count()
    if stored and not args.force_collection:
        raise SystemExit(
            f"Collection blocked: {DATABASE_PATH} already contains {stored} review(s). "
            "Use --force-collection only after verifying that a paid refresh is intended."
        )
    collector = ApifyGoogleReviewsCollector()
    started_at = datetime.now(timezone.utc).isoformat()
    if args.max_reviews is None:
        print('Collecting all available reviews per place...')
    else:
        print(f'Collecting up to {args.max_reviews} review(s) per place...')
    raw_items = collector.collect(urls=args.url, max_reviews=args.max_reviews, sort=args.sort, language=args.language, include_personal_data=args.personal_data)
    print(f'Apify returned {len(raw_items)} raw review item(s).')
    if not raw_items:
        print('No reviews were returned by Apify.')
        return
    places, reviews = normalize_items(raw_items)
    db = ReviewDatabase()
    try:
        db.upsert_places(places)
        inserted, updated = db.upsert_reviews(reviews)
        print(f'Places processed: {len(places)}')
        print(f'New reviews inserted: {inserted}')
        print(f'Existing reviews refreshed: {updated}')
        if args.export_csv or args.export_excel:
            EXPORT_DIR.mkdir(parents=True, exist_ok=True)
            for place in places:
                place_id = place['place_id']
                place_reviews = [review for review in reviews if review['place_id'] == place_id]
                export_rows = build_export_rows(place_reviews, place, include_raw_json=args.include_raw_json)
                safe_name = slugify(place.get('title') or place_id)
                csv_output, excel_output = export_paths(safe_name, len(places))
                if args.export_csv:
                    pd.DataFrame(export_rows).to_csv(csv_output, index=False, encoding='utf-8-sig')
                    print(f'CSV export: {csv_output}')
                if args.export_excel:
                    export_excel(excel_output, place, export_rows)
                    print(f'Excel export: {excel_output}')
        save_collection_summary(place_count=len(places), returned_reviews=len(reviews), inserted_reviews=inserted, updated_reviews=updated)
        write_collection_snapshot(args.url, started_at, reviews, collector)
    finally:
        db.close()

def places_command(_args):
    db = ReviewDatabase()
    try:
        rows = db.list_places()
        if not rows:
            print('No places stored yet.')
            return
        for row in rows:
            print(f"{row['title']} | {row['category']} | {row['city'] or '-'} | stored={row['stored_reviews']} | avg={row['stored_avg_rating']} | id={row['place_id']}")
        save_database_summary(rows)
    finally:
        db.close()

def stats_command(args):
    db = ReviewDatabase()
    try:
        summary, distribution = db.place_stats(args.place_id)
        if not summary:
            print('Place not found.')
            return
        print('\nPLACE SUMMARY')
        print('-' * 50)
        print(f"Name: {summary['title']}")
        print(f"Category: {summary['category']}")
        print(f"Location: {summary['city'] or '-'}, {summary['country_code'] or '-'}")
        print(f"Stored reviews: {summary['stored_reviews']}")
        print(f"Stored average: {summary['stored_avg_rating']}")
        print(f"Google overall rating: {summary['google_score']}")
        print(f"Google review count: {summary['google_review_count']}")
        print(f"Stored date range: {summary['oldest_review']} -> {summary['newest_review']}")
        print('\nRATING DISTRIBUTION')
        print('-' * 50)
        for row in distribution:
            print(f"{row['rating']} star: {row['count']}")
    finally:
        db.close()

def build_parser():
    parser = argparse.ArgumentParser(description='Generic place review intelligence collector.')
    sub = parser.add_subparsers(dest='command', required=True)
    collect = sub.add_parser('collect', help='Collect Google Maps reviews via Apify.')
    collect.add_argument('--url', action='append', required=True, help='Google Maps place URL. Repeat --url for multiple places.')
    collect.add_argument('--max-reviews', type=int, default=None, help='Maximum number of reviews per place. Omit to collect all available reviews.')
    collect.add_argument('--sort', default='newest', choices=['newest', 'mostRelevant', 'highestRanking', 'lowestRanking'])
    collect.add_argument('--language', default='en', help='Google Maps interface language. Reviews in all available languages are still collected.')
    collect.add_argument('--personal-data', action=argparse.BooleanOptionalAction, default=True, help='Collect reviewer Google profile data. Enabled by default. Use --no-personal-data to disable.')
    collect.add_argument('--export-csv', action='store_true', help='Export a readable CSV file.')
    collect.add_argument('--export-excel', action='store_true', help='Export a formatted Excel .xlsx file.')
    collect.add_argument('--include-raw-json', action='store_true', help='Include complete raw Apify JSON inside exported files.')
    collect.add_argument('--force-collection', action='store_true', help='Allow a verified paid refresh when reviews already exist.')
    collect.set_defaults(func=collect_command)
    places = sub.add_parser('places', help='List places stored in the database.')
    places.set_defaults(func=places_command)
    stats = sub.add_parser('stats', help='Show basic statistics for one stored place.')
    stats.add_argument('place_id')
    stats.set_defaults(func=stats_command)
    return parser

def main(argv: list[str] | None=None):
    parser = build_parser()
    if argv is None and len(sys.argv) == 1:
        argv = ['places']
    args = parser.parse_args(argv)
    args.func(args)
if __name__ == '__main__':
    main()
